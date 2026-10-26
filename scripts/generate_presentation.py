# scripts/generate_presentation.py
# Generates a professional 10-slide PowerPoint presentation (.pptx)
# formatted specifically according to academic presentation guidelines.

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333) # 16:9 widescreen
    prs.slide_height = Inches(7.5)

    # Theme Colors
    BG_COLOR = RGBColor(15, 23, 42)      # Deep slate #0f172a
    CARD_BG = RGBColor(30, 41, 59)       # Slate 800 #1e293b
    ACCENT_BLUE = RGBColor(56, 189, 248) # Sky blue #38bdf8
    ACCENT_GREEN = RGBColor(52, 211, 153)# Emerald #34d399
    TEXT_LIGHT = RGBColor(241, 245, 249) # Slate 100
    TEXT_MUTED = RGBColor(148, 163, 184) # Slate 400

    def add_blank_slide_with_bg():
        blank_slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_slide_layout)
        background = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        background.fill.solid()
        background.fill.fore_color.rgb = BG_COLOR
        background.line.color.rgb = BG_COLOR
        return slide

    def add_header(slide, title_text, category_tag=""):
        if category_tag:
            tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
            tf_tag = tag_box.text_frame
            tf_tag.word_wrap = True
            p_tag = tf_tag.paragraphs[0]
            p_tag.text = category_tag.upper()
            p_tag.font.size = Pt(11)
            p_tag.font.bold = True
            p_tag.font.color.rgb = ACCENT_BLUE

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT

    # ─── Slide 1: Title Slide ──────────────────────────────────────────────────
    s1 = add_blank_slide_with_bg()
    
    # Hero Title Box
    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(3.0))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    
    p_tag = tf1.paragraphs[0]
    p_tag.text = "ACADEMIC PROJECT PRESENTATION • SEMESTER 5"
    p_tag.font.size = Pt(12)
    p_tag.font.bold = True
    p_tag.font.color.rgb = ACCENT_BLUE
    
    p_title = tf1.add_paragraph()
    p_title.text = "Online Examination &\nProctoring System (ExamPortal)"
    p_title.font.size = Pt(40)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_LIGHT
    p_title.space_before = Pt(12)

    p_sub = tf1.add_paragraph()
    p_sub.text = "A Full-Stack Cloud Architecture for Remote Academic Assessments with Live Fraud Telemetry"
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = TEXT_MUTED
    p_sub.space_before = Pt(14)

    # Details card
    det_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(5.3), Inches(11.3), Inches(1.5))
    det_card.fill.solid()
    det_card.fill.fore_color.rgb = CARD_BG
    det_card.line.color.rgb = RGBColor(51, 65, 85)
    
    tf_det = det_card.text_frame
    tf_det.word_wrap = True
    p_det1 = tf_det.paragraphs[0]
    p_det1.text = "Presented by: Ankit Patel  |  Course: JavaScript Frameworks  |  Tech Stack: React 18, Node.js, Express, MongoDB Atlas, Socket.IO"
    p_det1.font.size = Pt(13)
    p_det1.font.bold = True
    p_det1.font.color.rgb = TEXT_LIGHT

    p_det2 = tf_det.add_paragraph()
    p_det2.text = "🌐 Live Public URL: https://online-examination-proctoring.vercel.app  •  Backend: Render Cloud  •  Status: 100% Deployed & Verified"
    p_det2.font.size = Pt(12)
    p_det2.font.color.rgb = ACCENT_GREEN
    p_det2.space_before = Pt(8)

    # ─── Slide 2: Definition, Objective & Problem Understanding (Point 1) ─────
    s2 = add_blank_slide_with_bg()
    add_header(s2, "1. Project Definition, Objective & Problem Understanding", "POINT 1 • CORE DEFINITION")

    # Left Column: Definition & Problem
    c1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = RGBColor(51, 65, 85)
    tf_c1 = c1.text_frame
    tf_c1.word_wrap = True
    
    p = tf_c1.paragraphs[0]
    p.text = "Problem Understanding"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    bullets_c1 = [
        ("Traditional Exams Bottlenecks: ", "Heavy printing and invigilation expenses, slow error-prone manual grading, geographical limitations."),
        ("Naive Quiz Portals Flaws: ", "Google Forms / basic LMS allow uncontrolled tab switching, AI assistant pasting (ChatGPT), and client-side clock tampering."),
        ("The Need: ", "A zero-plugin web platform providing server-authoritative timings, live cheat detection, and instant evaluation.")
    ]
    for b_title, b_desc in bullets_c1:
        p = tf_c1.add_paragraph()
        p.text = "• " + b_title + b_desc
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(12)

    # Right Column: Objectives
    c2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = RGBColor(51, 65, 85)
    tf_c2 = c2.text_frame
    tf_c2.word_wrap = True

    p = tf_c2.paragraphs[0]
    p.text = "Core Project Objectives"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    bullets_c2 = [
        ("Academic Integrity Protection: ", "Detect tab switching (Page Visibility API) and full-screen exits with live webcam preview."),
        ("Server-Authoritative Clock: ", "Calculate deadlines on backend (min(startedAt + duration, endTime)), preventing client-side timer manipulation."),
        ("Instant Auto-Grading: ", "Zero-latency evaluation of MCQs and True/False with configurable negative marking deductions."),
        ("Multi-Role Governance: ", "Dedicated portals for Admin, Teacher, and Student with immutable audit logging."),
        ("Cloud Production Availability: ", "Stateless JWT httpOnly cookie security deployed across Vercel, Render, and MongoDB Atlas.")
    ]
    for b_title, b_desc in bullets_c2:
        p = tf_c2.add_paragraph()
        p.text = "• " + b_title + b_desc
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(10)

    # ─── Slide 3: Technical Feasibility (Point 2) ──────────────────────────────
    s3 = add_blank_slide_with_bg()
    add_header(s3, "2. Technical Feasibility Study", "POINT 2 • FEASIBILITY ASSESSMENT")

    feasibility_data = [
        ("Frontend Tier", "React 18 + Vite 8 + Tailwind CSS v3", "Zero client installation; operates seamlessly on Chrome, Edge, Safari; lightweight bundle (< 600 KB gzip)."),
        ("Backend Tier", "Node.js 20 LTS + Express 5", "Asynchronous non-blocking event loop handles hundreds of concurrent heartbeat saves with low memory (~40MB)."),
        ("Database Tier", "MongoDB 7.0 + Mongoose 9", "Flexible JSON document schema ideal for polymorphic questions, options arrays, and nested student attempt answers."),
        ("Real-time Push", "Socket.IO 4 (WebSockets)", "Low-latency bidirectional room broadcasting (exam rooms & attempt rooms) for live watchdog alerts and warnings."),
        ("Security Stack", "JWT httpOnly Cookies + Helmet + Zod", "Complete defense against XSS (no localStorage tokens), NoSQL injection sanitizers, and CSRF protection."),
        ("Deployment Tier", "Vercel + Render + MongoDB Atlas", "Cost-effective, highly available, production-grade cloud infrastructure with automated HTTPS certificates.")
    ]

    for i, (f_title, f_tech, f_desc) in enumerate(feasibility_data):
        row = i // 2
        col = i % 2
        x = Inches(0.8 + col * 6.0)
        y = Inches(1.6 + row * 1.7)
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.6), Inches(1.5))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = RGBColor(51, 65, 85)
        tf = card.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f"{f_title}: {f_tech}"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE
        
        p2 = tf.add_paragraph()
        p2.text = f_desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(4)

    # ─── Slide 4: Suitable Process Model & Why (Point 3) ──────────────────────
    s4 = add_blank_slide_with_bg()
    add_header(s4, "3. Suitable Process Model: Agile Iterative & Incremental", "POINT 3 • METHODOLOGY JUSTIFICATION")

    # Left: Sprints
    card_s = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    card_s.fill.solid()
    card_s.fill.fore_color.rgb = CARD_BG
    card_s.line.color.rgb = RGBColor(51, 65, 85)
    tf_s = card_s.text_frame
    tf_s.word_wrap = True

    p = tf_s.paragraphs[0]
    p.text = "11-Phase Incremental Sprints"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    sprints = [
        "Phase 1: Architecture, scaffold & design system tokens",
        "Phase 2: Auth, JWT httpOnly cookies & protected routes",
        "Phase 3: Admin dashboard & platform user management",
        "Phase 4: Teacher exam creation & question bank palette",
        "Phase 5-6: Student exam UI, auto-saving & auto-grading",
        "Phase 7: Evaluation portal & faculty score analytics",
        "Phase 8-9: Live proctoring watchdog & audit logging",
        "Phase 10: Automated test suite (Jest + Vitest) & Error Boundary",
        "Phase 11: Security hardening (NoSQL sanitizer, rate limits) & Cloud Deploy"
    ]
    for s_item in sprints:
        p = tf_s.add_paragraph()
        p.text = "✔ " + s_item
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(6)

    # Right: Justification Why
    card_w = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2))
    card_w.fill.solid()
    card_w.fill.fore_color.rgb = CARD_BG
    card_w.line.color.rgb = RGBColor(51, 65, 85)
    tf_w = card_w.text_frame
    tf_w.word_wrap = True

    p = tf_w.paragraphs[0]
    p.text = "Why Agile Incremental Model?"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    reasons = [
        ("Risk De-escalation: ", "Developing core data models and auth before real-time event listeners and countdown timers ensured structural integrity."),
        ("Early Validation: ", "Automated grading calculations and server deadlines were tested in isolation with Jest before UI integration."),
        ("Evolving Proctoring Needs: ", "Proctoring expanded from basic tab-switch logging to bidirectional WebSockets, smoothly accommodated in late iterations."),
        ("Continuous Regression Testing: ", "Every sprint closed with automated verification (17 backend + 13 frontend tests passing).")
    ]
    for r_title, r_desc in reasons:
        p = tf_w.add_paragraph()
        p.text = "• " + r_title + r_desc
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(12)

    # ─── Slide 5: System Analysis - Use Case Diagram (Point 4) ────────────────
    s5 = add_blank_slide_with_bg()
    add_header(s5, "4. Software Analysis: Use Case Diagram & Actor Matrix", "POINT 4 • UML ANALYSIS")

    actors = [
        ("Admin Actor", ACCENT_BLUE, [
            "View platform metrics (users, exams, attempts, uptime)",
            "Filter users by role & toggle account active/inactive",
            "Review immutable security audit logs",
            "Monitor system health & readiness endpoints"
        ]),
        ("Teacher Actor", ACCENT_GREEN, [
            "Create & configure exams (duration, schedule, negative marks)",
            "Curate questions (MCQ Single, MCQ Multiple, True/False)",
            "Publish / Archive examinations",
            "Monitor live proctoring events & send direct student warnings",
            "Manually grade descriptive questions & release results",
            "Analyze pass rates & score distributions"
        ]),
        ("Student Actor", RGBColor(244, 114, 182), [
            "Authenticate / Register account",
            "Browse active published examinations",
            "Grant fullscreen & camera permissions",
            "Take exam with live countdown & auto-save heartbeat",
            "Trigger proctoring watchdog telemetry on tab switch",
            "Submit attempt & view released scorecard"
        ])
    ]

    for i, (a_name, a_col, a_tasks) in enumerate(actors):
        x = Inches(0.8 + i * 3.9)
        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.6), Inches(3.7), Inches(5.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = RGBColor(51, 65, 85)
        tf = card.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = a_name
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = a_col

        for t in a_tasks:
            p = tf.add_paragraph()
            p.text = "• " + t
            p.font.size = Pt(12)
            p.font.color.rgb = TEXT_LIGHT
            p.space_before = Pt(10)

    # ─── Slide 6: System Analysis - ER & Data Architecture (Point 4) ──────────
    s6 = add_blank_slide_with_bg()
    add_header(s6, "4. Software Analysis: Entity-Relationship & Database Architecture", "POINT 4 • DATA MODEL")

    collections = [
        ("Users Collection", "Stores Admin, Teacher, and Student accounts with bcrypt-hashed passwords (cost factor 12) and accountStatus flags."),
        ("Exams Collection", "Configures exam windows (startTime, endTime), durationMinutes, passingMarks, proctoring settings, and negative marking rules."),
        ("Questions Collection", "Polymorphic schema for MCQs and descriptive questions. Crucially stores answerKey with select: false for network security."),
        ("ExamAttempts Collection", "Tracks attempt lifecycle (IN_PROGRESS, SUBMITTED, EXPIRED), server computed deadlineAt, score, and embedded student answers."),
        ("ProctoringEvents Collection", "Logs integrity violations (TAB_SWITCH, FULLSCREEN_EXIT), durationMs, and teacher review status (UNREVIEWED, REVIEWED, DISMISSED)."),
        ("AuditLogs Collection", "Immutable append-only record capturing actorId, action, targetId, and timestamp for all privileged administrative events.")
    ]

    for i, (c_name, c_desc) in enumerate(collections):
        row = i // 2
        col = i % 2
        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + col * 6.0), Inches(1.6 + row * 1.7), Inches(5.6), Inches(1.5))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = RGBColor(51, 65, 85)
        tf = card.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = c_name
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE

        p2 = tf.add_paragraph()
        p2.text = c_desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_LIGHT
        p2.space_before = Pt(4)

    # ─── Slide 7: Student Exam & Live Proctoring Engine ───────────────────────
    s7 = add_blank_slide_with_bg()
    add_header(s7, "Student Examination & Real-Time Proctoring Watchdog", "INTEGRITY ENGINE")

    features = [
        ("Server-Authoritative Countdown", "The countdown timer is anchored strictly to database deadlineAt (min(startedAt + duration, endTime)). Client system clock tampering cannot add exam time."),
        ("HTML5 Fullscreen Enforcement", "Candidates are locked into full-screen mode. Pressing Esc or exiting triggers an immediate integrity warning dialog and database violation log."),
        ("Page Visibility & Tab Tracking", "Listens for document visibilitychange and window blur. Navigating away to search engines or AI tools immediately increments violation counter."),
        ("Incremental Heartbeat Auto-Save", "Option selections trigger debounced PUT /api/attempts/:id/save. If internet temporarily drops, progress is safely stored in local state and restored upon reconnection."),
        ("Bidirectional WebSockets", "Proctoring violations push instantly via Socket.IO room (exam:examId). Instructors can issue immediate pop-up warnings (student:warning) to candidate screen.")
    ]

    for i, (f_title, f_desc) in enumerate(features):
        card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6 + i * 1.05), Inches(11.7), Inches(0.95))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = RGBColor(51, 65, 85)
        tf = card.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"{i+1}. {f_title}: "
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = ACCENT_GREEN

        run = p.add_run()
        run.text = f_desc
        run.font.size = Pt(12)
        run.font.bold = False
        run.font.color.rgb = TEXT_LIGHT

    # ─── Slide 8: Automated Grading & Analytics ───────────────────────────────
    s8 = add_blank_slide_with_bg()
    add_header(s8, "Automated Grading Engine & Faculty Analytics", "ASSESSMENT MODULE")

    # Left: Grading Engine
    g_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    g_card.fill.solid()
    g_card.fill.fore_color.rgb = CARD_BG
    g_card.line.color.rgb = RGBColor(51, 65, 85)
    tf_g = g_card.text_frame
    tf_g.word_wrap = True

    p = tf_g.paragraphs[0]
    p.text = "Objective Auto-Grading Engine"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    g_points = [
        ("Zero-Latency Evaluation: ", "On POST /submit, single-choice, multi-choice, and True/False questions are instantly graded against question.answerKey."),
        ("Negative Marking Penalty: ", "Supports configurable penalty deductions (e.g., -0.25 penalty), eliminating guessing advantages. Overall score clamped to Math.max(0, score)."),
        ("Subjective Question Workflow: ", "Descriptive and Short Answer questions flag attempt as MANUAL_PENDING for faculty review."),
        ("Teacher Evaluation Modal: ", "Instructors award marks, write custom qualitative remarks, and finalize scorecards before releasing results.")
    ]
    for pt, pd in g_points:
        p = tf_g.add_paragraph()
        p.text = "• " + pt + pd
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(12)

    # Right: Analytics
    a_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2))
    a_card.fill.solid()
    a_card.fill.fore_color.rgb = CARD_BG
    a_card.line.color.rgb = RGBColor(51, 65, 85)
    tf_a = a_card.text_frame
    tf_a.word_wrap = True

    p = tf_a.paragraphs[0]
    p.text = "Faculty Analytics Dashboard"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    a_points = [
        ("Pass / Fail Rates: ", "Automated visual calculation of passing percentages based on configured exam passingMarks."),
        ("Class Score Distribution: ", "Displays highest score, lowest score, class average, and median performance metrics."),
        ("Question-Level Difficulty Index: ", "Breakdown of accuracy percentages per question to identify confusing topics or curriculum gaps."),
        ("Controlled Results Release: ", "Students cannot view answer keys or scores until the instructor explicitly clicks 'Release Results'.")
    ]
    for pt, pd in a_points:
        p = tf_a.add_paragraph()
        p.text = "• " + pt + pd
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(12)

    # ─── Slide 9: Security Hardening & Testing ─────────────────────────────────
    s9 = add_blank_slide_with_bg()
    add_header(s9, "Security Hardening & Quality Assurance", "VERIFICATION")

    sec_cards = [
        ("NoSQL Injection Sanitization", "Recursive middleware sanitizes req.body, req.query, and req.params, stripping any keys starting with $ or containing . (prevents operator injection)."),
        ("Dual Token httpOnly Security", "15-minute access tokens and 7-day refresh tokens stored in httpOnly, secure, sameSite: none cookies (immune to XSS token theft in localStorage)."),
        ("Multi-Tier Rate Limiting", "globalLimiter (200 req/min), authLimiter (15 req/15min), and proctoringLimiter (120 req/min) prevent brute-force and DoS attacks."),
        ("React Production Error Boundary", "Catches unexpected runtime UI errors without blank screens, providing user recovery options ('Reload Page' and 'Go to Home')."),
        ("17 Backend Tests (Jest + Supertest)", "100% pass rate: Covers Health Check, Auth lifecycle, Exam CRUD, Publishing, Auto-grading, and Proctoring events."),
        ("13 Frontend Tests (Vitest)", "100% pass rate: Covers Button states, Badge auto-coloring, Alert dismissal, ErrorBoundary catch, and timer formatters.")
    ]

    for i, (s_title, s_desc) in enumerate(sec_cards):
        row = i // 2
        col = i % 2
        card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + col * 6.0), Inches(1.6 + row * 1.7), Inches(5.6), Inches(1.5))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = RGBColor(51, 65, 85)
        tf = card.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = s_title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE

        p2 = tf.add_paragraph()
        p2.text = s_desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_LIGHT
        p2.space_before = Pt(4)

    # ─── Slide 10: Conclusion & Live Demonstration ────────────────────────────
    s10 = add_blank_slide_with_bg()
    add_header(s10, "Conclusion & Live Demonstration", "PROJECT DEMO")

    # Big Demo Card
    demo_card = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    demo_card.fill.solid()
    demo_card.fill.fore_color.rgb = CARD_BG
    demo_card.line.color.rgb = RGBColor(51, 65, 85)
    tf_demo = demo_card.text_frame
    tf_demo.word_wrap = True

    p = tf_demo.paragraphs[0]
    p.text = "🚀 Live Cloud Deployment Links & Demo Accounts"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    demo_items = [
        ("Frontend URL: ", "https://online-examination-proctoring.vercel.app  (Vercel Edge Global CDN)"),
        ("Backend API: ", "https://online-examination-proctoring.onrender.com  (Render Container Service)"),
        ("Cloud Database: ", "MongoDB Atlas M0 Cluster (cluster0.4hrusic.mongodb.net)"),
        ("Admin Login: ", "admin@examportal.dev  |  Password: Admin@1234!  (Inspect metrics & audit trail)"),
        ("Teacher Login: ", "prof.ahmad@college.edu  |  Password: Teacher@1234!  (Exam creation & live watchdog review)"),
        ("Student Login: ", "jane.student@examportal.dev  |  Password: Student@1234!  (Take exam with webcam & timer)"),
        ("Project Status: ", "All 11 implementation phases complete, 30 automated tests passing, 100% production ready.")
    ]

    for d_title, d_val in demo_items:
        p = tf_demo.add_paragraph()
        p.text = "• " + d_title + d_val
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(10)

    # Save presentation
    output_filename = "Online_Examination_Proctoring_Presentation.pptx"
    prs.save(output_filename)
    print(f"Presentation saved successfully as '{output_filename}'")

if __name__ == '__main__':
    create_presentation()
